"""把一批真实表单脱敏成可以在本机学习、评测的数据（CLAUDE.md §11）。

    uv run python ai/eval/desensitize.py real_data/<学年>/原始 real_data/<学年>/脱敏 \\
        --private real_data/<学年>/私密

- 支持 .xlsx、.pdf（带文字层），以及 .doc/.docx/.wps（先用 LibreOffice 转成 PDF）。
- 每个源文件输出一个 JSON：{"source", "tables": [[行...]...], "text"}，姓名、学号都已替换。
- 对照表和密钥写在 --private 目录；该目录和输出目录都必须在 real_data/ 或 data/ 下（不入库）。
- 只在终端打印统计（文件数、人数、泄漏数、可疑数），不打印任何内容。
"""

import argparse
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

from zongce_ai.desensitize import (
    Desensitizer,
    Document,
    Table,
    collect,
    column_residue,
    suspicious,
)
from zongce_ai.names import majors_from

SAFE_ROOTS = ("real_data", "data")


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\n", " ").strip()


def read_xlsx(path: Path, name: str) -> Document:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    doc = Document(name)
    for sheet in workbook.worksheets:
        rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        doc.tables.append(Table([r for r in rows if any(r)]))
    workbook.close()
    return doc


def read_pdf(path: Path, name: str) -> Document:  # pragma: no cover - 依赖 pdfplumber 与真实 PDF
    import pdfplumber

    doc = Document(name)
    texts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                doc.tables.append(Table([[_cell(c) for c in row] for row in table]))
            texts.append(page.extract_text() or "")
    doc.text = "\n".join(texts)
    return doc


def read_office(path: Path, name: str) -> Document:  # pragma: no cover - 依赖 LibreOffice
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(path)],
            capture_output=True,
            timeout=180,
            check=False,
        )
        converted = Path(tmp) / (path.stem + ".pdf")
        if not converted.exists():
            raise RuntimeError("LibreOffice 转换失败")
        return read_pdf(converted, name)


READERS = {".xlsx": read_xlsx, ".pdf": read_pdf, ".doc": read_office, ".docx": read_office,
           ".wps": read_office}  # fmt: skip


def _check_safe(path: Path) -> None:
    if not any(part in SAFE_ROOTS for part in path.resolve().parts):
        sys.exit(f"{path} 不在 real_data/ 或 data/ 下；脱敏数据和对照表不能放在会被提交的位置")


def run(source: Path, output: Path, private: Path) -> dict[str, Any]:
    _check_safe(output)
    _check_safe(private)
    documents: list[Document] = []
    failed: list[str] = []
    for path in sorted(p for p in source.rglob("*") if p.is_file()):
        reader = READERS.get(path.suffix.lower())
        if reader is None:
            continue
        name = path.relative_to(source).as_posix()
        try:
            documents.append(reader(path, name))
        except Exception as exc:  # 单个文件读失败不影响其他文件  # pragma: no cover
            failed.append(f"{name}：{type(exc).__name__}")
    people = collect(documents)
    private.mkdir(parents=True, exist_ok=True)
    key = private / "secret.key"
    if not key.exists():
        key.write_bytes(os.urandom(32))
    tool = Desensitizer(people, key.read_bytes())
    (private / "对照表.json").write_text(
        json.dumps({"names": tool.mapping, "ids": tool.id_mapping}, ensure_ascii=False, indent=1),
        encoding="utf-8",
    )
    majors = majors_from(people.class_codes)
    files = []
    for doc in documents:
        tables = [[[tool.scrub(c) for c in row] for row in t.rows] for t in doc.tables]
        text = tool.scrub(doc.text)
        flat = "\n".join(["\n".join(" ".join(r) for r in t) for t in tables] + [text])
        target = output / (doc.name + ".json")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            json.dumps({"source": doc.name, "tables": tables, "text": text}, ensure_ascii=False),
            encoding="utf-8",
        )
        files.append(
            {
                "file": doc.name,
                "tables": len(tables),
                "leaks": tool.leaks(flat) + column_residue(tables),
                "suspicious": suspicious(flat, majors),
            }
        )
    return {
        "files": files,
        "failed": failed,
        "students": len(people.students),
        "teachers": len(people.teachers),
        "ids": len(people.ids),
        "leaks": sum(f["leaks"] for f in files),
        "suspicious": sum(f["suspicious"] for f in files),
    }


def main(argv: list[str] | None = None) -> dict[str, Any]:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawTextHelpFormatter
    )
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--private", type=Path, required=True, help="对照表与密钥目录")
    args = parser.parse_args(argv)
    summary = run(args.source, args.output, args.private)
    print(json.dumps(summary, ensure_ascii=False, indent=1))
    return summary


if __name__ == "__main__":
    main()
