"""真实表单脱敏：收集姓名与学号、一致替换、泄漏检查（只用虚构数据）。"""

import importlib.util
import json
from pathlib import Path
from types import ModuleType

import pytest
from openpyxl import Workbook

from ai_helpers import REPO_ROOT
from zongce_ai.desensitize import (
    Desensitizer,
    Document,
    People,
    Table,
    collect,
    column_residue,
    refine,
    suspicious,
)
from zongce_ai.names import DEFAULT_MAJORS

TEXT = """团委
活动名称：某主题团日活动
一等奖：2 分
计科 2501 张三 网合25w3李四 测绘2401王 五
二等奖：1.5 分
刘 一 陈小明 周九
指导教师：吴老大、郑老二
联系电话 13800000000
"""


def _documents() -> list[Document]:
    roster = Table(
        [
            ["ID", "姓名", "学号", "班级", "创新实践课程学分-积分", "2025年下半年志愿汇时长"],
            ["12345678", "张三", "202500000001", "计科2501", "3.5", "1"],
            ["12345679", "赵六", "202500000002", "测绘2401", "1", "0"],
        ]
    )
    awards = Table(
        [
            ["序号", "赛事名称", "获奖情况", "获奖学生总人名"],
            ["1", "参赛作品名称很长的某某全国大学生竞赛", "省三等奖", "钱七,孙八"],
            ["2", "某竞赛", "国一等奖", "张 三"],
        ]
    )
    projects = Table(
        [
            ["项目名称", "项目负责人", "项目成员（姓名/学号）", "指导教师"],
            ["某项目", "李四", "李四/202500000003, 周九/202500000004", "吴老大"],
        ]
    )
    dorm = Table([["宿舍", "班级", "名单"], ["学1-101", "计科2501", "陈小明 刘一 王五"]])
    return [
        Document("roster.xlsx", tables=[roster]),
        Document("awards.pdf", tables=[awards, projects, dorm], text=TEXT),
    ]


def test_collect_people() -> None:
    people = collect(_documents())
    assert people.students == {"张三", "李四", "王五", "赵六", "钱七", "孙八", "刘一", "陈小明",
                               "周九"}  # fmt: skip
    assert people.teachers == {"吴老大", "郑老二"}
    assert {"202500000001", "202500000004", "12345678", "13800000000"} <= people.ids
    assert "测绘2401" in people.class_codes


def test_header_carries_over_to_next_page() -> None:
    page1 = Table([["序号", "赛事名称", "获奖学生总人名"], ["1", "某赛", "甲乙"]])
    page2 = Table([["2", "某赛", "丙丁, 戊己"]])  # 第二页没有表头
    other = Table([["1", "2"], ["庚辛", "壬癸"]])  # 列数不同，不沿用
    people = collect([Document("x.pdf", tables=[page1, page2, other])])
    assert people.students == {"甲乙", "丙丁", "戊己"}
    assert column_residue([page1.rows, page2.rows]) == 2  # 按格计数


def test_refine_finds_names_next_to_known_ones() -> None:
    roster = Table(
        [["姓名", "班级"], ["张三", "计科2501"], ["李四", "计科2501"], ["王五", "计科2501"]]
    )
    text = (
        "4*100米：张三 赵小六 李四\n"
        "张三，钱七，李四\n"
        "张三孙八\n"  # 粘在一起
        "学23-108-周九 学23-109-吴十\n"
        "留学生 张三 BLAKE JAN\n"
        "AI APP 大赛 张三\n"  # 缩写不是姓名
        "指导教师 张三 李四\n"  # “教师”是代号的前缀，不能当成姓名
        "一等奖 2 分\n"
    )
    docs = [Document("roster.xlsx", tables=[roster]), Document("x.pdf", text=text)]
    one_round = refine(collect(docs), docs, b"k" * 32, rounds=1)
    assert "赵小六" in one_round.students
    people = refine(collect(docs), docs, b"k" * 32)
    assert {"赵小六", "钱七", "孙八", "周九", "吴十", "BLAKEJAN"} <= people.students
    assert "AIAPP" not in people.students
    assert "教师" not in people.students
    tool = Desensitizer(people, b"k" * 32)
    scrubbed = tool.scrub(text)
    for real in ("赵小六", "钱七", "孙八", "周九", "BLAKE JAN"):
        assert real not in scrubbed
    assert "一等奖" in scrubbed
    assert column_residue([[["姓名"], ["一等奖：2 分"]]]) == 0


def test_guessed_words_that_are_not_names() -> None:
    doc = Document(
        "三查汇总+优秀宿舍.pdf", text="优秀宿舍 宿舍\n计科2501 一等奖\n获奖 名单\n张三 李四"
    )
    assert collect([doc]).students == {"张三", "李四"}


def test_scrub_is_consistent_and_leak_free() -> None:
    documents = _documents()
    tool = Desensitizer(collect(documents), b"k" * 32)
    scrubbed = tool.scrub(TEXT)
    for real in ("张三", "李四", "王 五", "陈小明", "吴老大", "13800000000"):
        assert real not in scrubbed
    assert tool.leaks(scrubbed) == 0
    assert tool.leaks(TEXT) > 0
    zhang = tool.mapping["张三"]
    assert zhang.startswith("学生")
    assert tool.scrub("张 三") == zhang  # 同一人的不同写法换成同一代号
    assert tool.mapping["吴老大"].startswith("教师")
    fake = tool.scrub("202500000001")
    assert len(fake) == 12
    assert fake != "202500000001"
    assert fake == tool.scrub("202500000001")
    assert tool.scrub("99999999") == "00000000"  # 未收集到的长数字也不保留
    assert Desensitizer(collect(documents), b"k" * 32).mapping == tool.mapping  # 同一密钥可复现
    assert suspicious(scrubbed, (*DEFAULT_MAJORS, "测绘")) == 0
    assert suspicious("计科2501 无名氏\n甲乙 丙丁", DEFAULT_MAJORS) == 2


def test_column_residue() -> None:
    tables = [
        [["姓名", "指导教师", "班级"], ["学生0001", "教师001", "计科2501"], ["张三", "无", ""]]
    ]
    assert column_residue(tables) == 1
    assert column_residue([[["姓名"], ["学生0001、学生0002"], ["智能 2301 学生0003"]]]) == 0


def test_empty_people() -> None:
    tool = Desensitizer(People(), b"k" * 32)
    assert tool.scrub("没有人名 12345678") == "没有人名 00000000"
    assert tool.leaks("没有人名") == 0


def _script() -> ModuleType:
    spec = importlib.util.spec_from_file_location(
        "desensitize", REPO_ROOT / "ai" / "eval" / "desensitize.py"
    )
    assert spec is not None
    assert spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_script_on_xlsx(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    source = tmp_path / "原始"
    source.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in (
        ["姓名", "学号", "班级"],
        ["张三", 202500000001, "计科2501"],
        ["李四", 2.025e11, None],
    ):
        sheet.append(row)
    workbook.save(source / "名单.xlsx")
    (source / "说明.txt").write_text("忽略", encoding="utf-8")
    out, private = tmp_path / "real_data" / "脱敏", tmp_path / "real_data" / "私密"
    summary = _script().main([str(source), str(out), "--private", str(private)])
    assert (summary["students"], summary["leaks"], summary["failed"]) == (2, 0, [])
    data = json.loads((out / "名单.xlsx.json").read_text(encoding="utf-8"))
    assert data["tables"][0][0] == ["姓名", "学号", "班级"]
    assert data["tables"][0][1][0].startswith("学生")
    assert "202500000001" not in json.dumps(data)
    assert (private / "secret.key").exists()
    assert "张三" not in capsys.readouterr().out  # 终端只打印统计
    mapping = json.loads((private / "对照表.json").read_text(encoding="utf-8"))
    assert set(mapping["names"]) == {"张三", "李四"}


def test_script_refuses_unsafe_output(tmp_path: Path) -> None:
    with pytest.raises(SystemExit, match="real_data"):
        _script().main([str(tmp_path), str(tmp_path / "out"), "--private", str(tmp_path / "p")])
