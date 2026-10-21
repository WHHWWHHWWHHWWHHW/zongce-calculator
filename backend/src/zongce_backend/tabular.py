"""读取上传的表格（.xlsx / .csv）为字符串二维表。"""

import csv
import io
from decimal import Decimal

from openpyxl import load_workbook

from .errors import AppError

MAX_ROWS = 20000


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        # Excel 以二进制浮点存数字；repr 给出能还原该值的最短十进制串（如 3.6524）
        text = repr(value)
        return str(int(value)) if value.is_integer() else format(Decimal(text), "f")
    return str(value).strip()


def read_table(filename: str, data: bytes) -> list[list[str]]:
    """读取第一个工作表（CSV 只有一个表）。"""
    return read_sheets(filename, data, first_only=True)[0][1]


def read_sheets(
    filename: str, data: bytes, *, first_only: bool = False
) -> list[tuple[str, list[list[str]]]]:
    """读取全部工作表：[(工作表名, 行)]，已去掉空行。CSV 的工作表名为空字符串。"""
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl 对损坏文件抛出多种异常
            raise AppError("无法读取 Excel 文件，请确认是 .xlsx 格式") from exc
        sheets = workbook.worksheets[:1] if first_only else workbook.worksheets
        tables = [
            (
                sheet.title,
                _clean([[_cell_text(v) for v in r] for r in sheet.iter_rows(values_only=True)]),
            )
            for sheet in sheets
        ]
        workbook.close()
        return tables
    if name.endswith(".csv"):
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:  # pragma: no cover - gb18030 几乎能解码任意字节
            raise AppError("无法识别 CSV 文件编码，请另存为 UTF-8")
        return [("", _clean([[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]))]
    raise AppError("只支持 .xlsx 或 .csv 文件")


def _clean(rows: list[list[str]]) -> list[list[str]]:
    rows = [row for row in rows if any(cell for cell in row)]
    if len(rows) > MAX_ROWS:
        raise AppError(f"表格行数超过上限 {MAX_ROWS}")
    return rows
