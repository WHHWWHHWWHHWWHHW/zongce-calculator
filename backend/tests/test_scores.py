"""学院成绩导入、预估、核算（端到端复现 G1）、排名与导出。"""

import io
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend_helpers import API, World, upload, xlsx
from zongce_backend.models import AuditLog

G1_ROW = [3.6524, 2.5, 4, 85, 92, "9, 8.5, 10, 7, 9.5, 8, 9, 10, 6.5, 8", 9.5, 4.5, 5, 38]
HEADER = ["学号", "学分绩点", "第二课堂", "公益劳动", "体测成绩", "宿舍成绩", "政治素养",
          "品德修养", "纪律观念", "诚信评价", "能力分"]  # fmt: skip


def import_g1(client: TestClient, world: World, *usernames: str) -> Any:
    data = xlsx([HEADER, *[[u, *G1_ROW] for u in usernames]])
    url = f"{API}/years/{world.year_id}/college-data/import"
    return client.post(url, files=upload("学院成绩.xlsx", data), headers=world.admin)


def approve_ability_38(client: TestClient, world: World, headers: dict[str, str]) -> None:
    for code, option in (
        ("honor_national", "national"),
        ("honor_provincial", "provincial"),
        ("honor_university", "top_ten"),
        ("honor_university", "excellent"),
    ):
        body = {
            "year_id": world.year_id,
            "title": "测试",
            "catalog_code": code,
            "option_code": option,
        }
        did = client.post(f"{API}/declarations", json=body, headers=headers).json()["id"]
        client.post(f"{API}/declarations/{did}/submit", headers=headers)
        review = {"action": "approve", "reason": "测试"}
        client.post(f"{API}/declarations/{did}/review", json=review, headers=world.admin)


def compute(client: TestClient, world: World, **body: Any) -> Any:
    url = f"{API}/years/{world.year_id}/scores/compute"
    return client.post(url, json=body, headers=world.admin).json()


def test_import_mapping(client: TestClient, world: World, db: Session) -> None:
    result = import_g1(client, world, "TEST0002", "TEST9999", "TESTADMIN")
    body = result.json()
    assert body["mapped_columns"]["学分绩点"] == "gpa"
    assert body["mapped_columns"]["政治素养"] == "political"
    assert body["ignored_columns"] == ["能力分"]
    assert [e["row"] for e in body["errors"]] == [3, 4]
    assert (body["changed"], body["unchanged"]) == (9, 0)
    again = import_g1(client, world, "TEST0002").json()
    assert (again["changed"], again["unchanged"]) == (0, 9)
    records = client.get(
        f"{API}/years/{world.year_id}/college-data", params={"username": "TEST0002"},
        headers=world.admin,
    ).json()  # fmt: skip
    by_field = {r["field"]: r for r in records}
    assert by_field["gpa"]["value"] == "3.6524"
    assert by_field["political"]["ratings"] == [
        "9",
        "8.5",
        "10",
        "7",
        "9.5",
        "8",
        "9",
        "10",
        "6.5",
        "8",
    ]
    assert by_field["morality"]["value"] == "9.5"
    logs = db.scalars(select(AuditLog).where(AuditLog.action == "college_record.changed")).all()
    assert len(logs) == 9
    assert logs[0].reason is not None
    assert "学院成绩.xlsx" in logs[0].reason


def test_import_errors_and_manual_update(client: TestClient, world: World) -> None:
    url = f"{API}/years/{world.year_id}/college-data/import"

    def post(rows: list[list[Any]]) -> Any:
        return client.post(url, files=upload("x.xlsx", xlsx(rows)), headers=world.admin).json()

    assert post([["姓名", "绩点"], ["a", 1]])["error"]["message"] == "表头缺少“学号”列"
    assert (
        post([["学号", "备注"], ["TEST0002", "x"]])["error"]["message"]
        == "没有识别出任何成绩列，请检查表头"
    )
    bad = post([["学号", "体测", "体测成绩"], ["TEST0002", "缺考", ""]])
    assert bad["ignored_columns"] == ["体测成绩"]
    assert bad["errors"][0]["message"] == "“体测”无法识别数值"

    edit = f"{API}/years/{world.year_id}/college-data/TEST0002"
    value = client.put(
        f"{edit}/fitness_test", json={"value": "72", "reason": "补测"}, headers=world.admin
    )
    assert value.json()["value"] == "72"
    ratings = client.put(
        f"{edit}/integrity", json={"ratings": ["5", "4"], "reason": "补录"}, headers=world.admin
    )
    assert ratings.json()["ratings"] == ["5", "4"]
    wrong = client.put(
        f"{edit}/gpa", json={"ratings": ["5", "4"], "reason": "x"}, headers=world.admin
    )
    assert wrong.status_code == 400
    both = client.put(f"{edit}/gpa", json={"reason": "x"}, headers=world.admin)
    assert both.status_code == 422
    missing = client.put(
        f"{API}/years/{world.year_id}/college-data/NOBODY/gpa", json={"value": "1", "reason": "x"},
        headers=world.admin,
    )  # fmt: skip
    assert missing.status_code == 404
    # 学生只能看到自己的，审核员看到本班的
    own = client.get(f"{API}/years/{world.year_id}/college-data", headers=world.s2).json()
    assert own == []
    seen = client.get(f"{API}/years/{world.year_id}/college-data", headers=world.reviewer).json()
    assert {r["student_username"] for r in seen} == {"TEST0002"}


def test_preview(client: TestClient, world: World) -> None:
    body = {
        "gpa": "3.6524",
        "political": {"ratings_text": "9,8.5,10,7,9.5,8,9,10,6.5,8"},
        "morality": {"score": "9.5"},
        "discipline": {"ratings": ["4.5"]},
        "integrity": {"score": "5"},
        "second_classroom": "2.5",
        "public_service": "4",
        "fitness_test": "85",
        "dorm_check": "92",
        "ability": "38",
    }
    result = client.post(f"{API}/score/preview", json=body, headers=world.s1).json()
    assert (result["final_score"], result["moral_total"], result["competence_total"]) == (
        "3.7252", "27.6250", "51.2500",
    )  # fmt: skip
    political = next(i for i in result["items"] if i["rule_id"] == "moral.political")
    assert (political["peer_count"], political["peer_trimmed_per_side"]) == (10, 1)
    assert (political["raw"], political["raw_display"]) == ("8.625", "8.6250")
    final = next(i for i in result["items"] if i["rule_id"] == "final")
    assert (final["raw"], final["raw_display"], final["score_display"]) == (
        "3.7252375", "3.7252", "3.7252",
    )  # fmt: skip
    empty = next(i for i in result["items"] if i["rule_id"] == "basic.fitness_test")
    assert (empty["raw"], empty["raw_display"]) == ("85", "85.0000")

    direct = {"quality_mode": "direct", "quality_direct": "80", "gpa": "4",
              "rule_version_id": world.rule_version_id}  # fmt: skip
    out = client.post(f"{API}/score/preview", json=direct, headers=world.s1).json()
    assert (out["final_score"], out["moral_total"]) == ("4.0000", None)
    float_input = client.post(f"{API}/score/preview", json={"gpa": 3.5}, headers=world.s1)
    assert float_input.status_code == 422  # 分数必须是字符串
    two = client.post(
        f"{API}/score/preview", json={"political": {"score": "1", "ratings": ["1"]}},
        headers=world.s1,
    )  # fmt: skip
    assert two.status_code == 422
    unknown = client.post(f"{API}/score/preview", json={"rule_version_id": 999}, headers=world.s1)
    assert unknown.status_code == 404


def test_preview_without_rules(client: TestClient, db: Session) -> None:
    from backend_helpers import auth, make_user
    from zongce_backend.models import Role

    make_user(db, "TEST0500", Role.STUDENT)
    response = client.post(f"{API}/score/preview", json={}, headers=auth(client, "TEST0500"))
    assert response.json()["error"]["message"] == "尚未导入任何规则版本"


def test_compute_end_to_end_g1(client: TestClient, world: World, db: Session) -> None:
    import_g1(client, world, "TEST0002", "TEST0003")
    approve_ability_38(client, world, world.s1)
    approve_ability_38(client, world, world.s2)
    first = compute(client, world)
    assert first == {"computed": 4, "changed": 4}  # 4 个学生账号（管理员不是学生）
    scores = client.get(f"{API}/years/{world.year_id}/scores", headers=world.s1).json()
    assert len(scores) == 1
    mine = scores[0]
    assert (mine["final_score"], mine["moral_total"], mine["competence_total"]) == (
        "3.7252", "27.6250", "51.2500",
    )  # fmt: skip
    assert mine["ability_total"] == "38.0000"
    assert mine["detail"]["rules"] == {"id": "ujn-v1", "version": "1.0.0"}
    political = next(i for i in mine["detail"]["items"] if i["rule_id"] == "moral.political")
    assert (political["raw"], political["raw_display"]) == ("8.625", "8.6250")
    assert len(mine["detail"]["ability"]["declarations"]) == 4

    assert compute(client, world, reason="复核") == {"computed": 4, "changed": 0}
    # 改一项数据后重新核算：只有该学生的综测成绩变化，并写审计日志
    client.put(
        f"{API}/years/{world.year_id}/college-data/TEST0003/fitness_test",
        json={"value": "59", "reason": "成绩更正"}, headers=world.admin,
    )  # fmt: skip
    assert compute(client, world, class_name="测试A班", reason="复核") == {
        "computed": 3,
        "changed": 1,
    }
    logs = db.scalars(
        select(AuditLog).where(AuditLog.action == "score.computed").order_by(AuditLog.id)
    ).all()
    assert (logs[-1].before, logs[-1].after, logs[-1].reason) == ("3.7252", "3.6627", "复核")

    rankings = client.get(f"{API}/years/{world.year_id}/rankings", headers=world.reviewer).json()
    assert [(r["class_name"], r["rank"], r["username"]) for r in rankings] == [
        ("测试A班", 1, "TEST0002"),
        ("测试A班", 2, "TEST0003"),
        ("测试A班", 3, "TEST0001"),
    ]
    everyone = client.get(f"{API}/years/{world.year_id}/rankings", headers=world.admin).json()
    assert [r["class_name"] for r in everyone].count("测试B班") == 1
    assert client.get(
        f"{API}/years/{world.year_id}/rankings", params={"class_name": "测试B班"},
        headers=world.reviewer,
    ).status_code == 403  # fmt: skip
    assert client.get(f"{API}/years/{world.year_id}/rankings", headers=world.s1).status_code == 403


def test_ties_share_rank_and_export(client: TestClient, world: World, db: Session) -> None:
    from zongce_backend.models import User

    student = db.scalar(select(User).where(User.username == "TEST0003"))
    assert student is not None
    student.name = '=HYPERLINK("http://example.invalid")'
    db.commit()
    import_g1(client, world, "TEST0002", "TEST0003")
    compute(client, world)
    rows = client.get(
        f"{API}/years/{world.year_id}/rankings", params={"class_name": "测试A班"},
        headers=world.admin,
    ).json()  # fmt: skip
    assert [r["rank"] for r in rows] == [1, 1, 3]
    exported = client.get(
        f"{API}/years/{world.year_id}/rankings/export", params={"class_name": "测试A班"},
        headers=world.admin,
    )  # fmt: skip
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats")
    sheet = load_workbook(io.BytesIO(exported.content)).worksheets[0]
    values = list(sheet.iter_rows(values_only=True))
    assert values[0][0] == "班级"
    assert values[0][-1] == "综合测评成绩"
    assert [v[1] for v in values[1:]] == [1, 1, 3]
    # 无能力分：素质 27.625 + 13.25 = 40.875；40.875÷20×25% + 3.6524×75% = 3.2502375
    assert str(values[1][-1]) == "3.2502"
    names = {sheet.cell(row=r, column=4).value: sheet.cell(row=r, column=4).data_type
             for r in range(2, sheet.max_row + 1)}  # fmt: skip
    assert names['=HYPERLINK("http://example.invalid")'] == "s"  # 以文本保存，不是公式
    reviewer_export = client.get(
        f"{API}/years/{world.year_id}/rankings/export", headers=world.reviewer
    )
    assert reviewer_export.status_code == 200
